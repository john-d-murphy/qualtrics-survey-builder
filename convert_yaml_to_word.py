#!/usr/bin/env python3
"""
Convert a YAML survey spec to a formatted Word document.

Usage:
    python convert_yaml_to_word.py --yaml survey.yaml --out survey_questions.docx
"""
import argparse
import re
import sys
from typing import Any, Dict, List, Optional

import yaml
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def strip_html(html: str) -> str:
    """Strip HTML tags and collapse whitespace, preserving line breaks for block elements."""
    # Replace block-level tags with newlines
    text = re.sub(r"</(p|h[1-6]|div|li)>", "\n", html, flags=re.IGNORECASE)
    text = re.sub(r"<br\s*/?\s*>", "\n", text, flags=re.IGNORECASE)
    # Remove all remaining tags
    text = re.sub(r"<[^>]+>", "", text)
    # Decode HTML entities
    text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    text = text.replace("&rsquo;", "\u2019").replace("&nbsp;", " ")
    # Collapse multiple whitespace on each line, strip blank lines
    lines = [" ".join(line.split()) for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def logic_description(q: Dict[str, Any]) -> Optional[str]:
    """Build a human-readable description of display logic."""
    dl = q.get("display_logic", {}).get("condition", {})
    if not dl:
        return None
    ref = dl.get("question", "?")
    if "selected" in dl:
        return f'[Show if "{ref}" = "{dl["selected"]}"]'
    if "is_not" in dl:
        return f'[Show if "{ref}" \u2260 "{dl["is_not"]}"]'
    return None


def add_question(doc, q: Dict[str, Any], num: int, block_title: str) -> None:
    """Add a single question to the document."""
    qtype = q.get("type", "")
    text = q.get("text", q.get("title", ""))
    qid = q.get("id", "")

    # Question number + text
    p = doc.add_paragraph()
    run = p.add_run(f"{num}. ")
    run.bold = True
    run.font.size = Pt(11)
    run = p.add_run(text)
    run.font.size = Pt(11)

    # Type annotation
    type_label = qtype.replace("_", " ").title()
    extra = ""
    if qtype == "multi_select":
        mn = q.get("min_choices")
        mx = q.get("max_choices")
        if mn and mx:
            extra = f" (select {mn}\u2013{mx})"
        elif mx:
            extra = f" (select up to {mx})"
        else:
            extra = " (select all that apply)"
    elif qtype == "single_select":
        extra = " (choose one)"
    elif qtype == "dropdown":
        extra = " (dropdown)"
    elif qtype == "matrix":
        extra = " (matrix)"
    elif qtype in ("bipolar_group", "bipolar_matrix"):
        pts = q.get("points", 7)
        extra = f" ({pts}-point scale per item)"
    elif qtype == "slider_group":
        extra = " (slider per item)"
    elif qtype == "text_entry":
        sub = q.get("subtype", "multi_line")
        extra = " (open text)" if sub == "multi_line" else " (single line)"
    elif qtype == "descriptive":
        extra = " (info text)"

    p2 = doc.add_paragraph()
    run = p2.add_run(f"   Type: {type_label}{extra}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Display logic
    logic = logic_description(q)
    if logic:
        p3 = doc.add_paragraph()
        run = p3.add_run(f"   {logic}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

    # Screen out
    so = q.get("screen_out")
    if so:
        cond = so.get("condition", {})
        sel = cond.get("selected", cond.get("is_not", ""))
        msg = so.get("message", "End survey")
        p4 = doc.add_paragraph()
        run = p4.add_run(f'   \u26a0 Screen-out if "{sel}" \u2192 {msg}')
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    # Choices / options
    if qtype in ("single_select", "multi_select", "dropdown"):
        choices = q.get("choices", [])
        for c in choices:
            bp = doc.add_paragraph(style="List Bullet")
            run = bp.add_run(c)
            run.font.size = Pt(10)
        if q.get("allow_other"):
            bp = doc.add_paragraph(style="List Bullet")
            run = bp.add_run("Other (please specify)")
            run.font.size = Pt(10)
            run.italic = True

    elif qtype == "matrix":
        rows = q.get("rows", [])
        columns = q.get("columns", [])
        # Add a small table
        table = doc.add_table(rows=len(rows) + 1, cols=len(columns) + 1)
        table.style = "Light Grid Accent 1"
        # Header row
        table.cell(0, 0).text = ""
        for j, col in enumerate(columns):
            cell = table.cell(0, j + 1)
            cell.text = col
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.bold = True
        # Data rows
        for i, row in enumerate(rows):
            cell = table.cell(i + 1, 0)
            cell.text = row
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
            for j in range(len(columns)):
                table.cell(i + 1, j + 1).text = "\u25cb"  # empty circle

    elif qtype in ("bipolar_group", "bipolar_matrix"):
        pairs = q.get("pairs", [])
        pts = q.get("points", 7)
        # Table: left | o o o o o o o | right
        table = doc.add_table(rows=len(pairs), cols=3)
        table.style = "Light Grid Accent 1"
        for i, pair in enumerate(pairs):
            left, right = pair
            cell_l = table.cell(i, 0)
            cell_l.text = left
            for p in cell_l.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.size = Pt(9)
            cell_m = table.cell(i, 1)
            cell_m.text = " ".join(["\u25cb"] * pts)
            for p in cell_m.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9)
            cell_r = table.cell(i, 2)
            cell_r.text = right
            for p in cell_r.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    elif qtype == "slider_group":
        items = q.get("items", [])
        labels = q.get("labels", {})
        left = labels.get("left", "Low")
        right = labels.get("right", "High")
        for item in items:
            bp = doc.add_paragraph(style="List Bullet")
            run = bp.add_run(
                f"{item}:  {left} \u2500\u2500\u2500\u2500\u2500 {right}"
            )
            run.font.size = Pt(10)

    elif qtype == "text_entry":
        p5 = doc.add_paragraph()
        run = p5.add_run("   [___________________________________]")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # Force response note
    if q.get("force_response"):
        p6 = doc.add_paragraph()
        run = p6.add_run("   * Required")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)


def main():
    ap = argparse.ArgumentParser(
        description="Convert YAML survey spec to Word document."
    )
    ap.add_argument("--yaml", required=True, help="YAML survey spec path")
    ap.add_argument("--out", required=True, help="Output .docx path")
    args = ap.parse_args()

    with open(args.yaml, "r", encoding="utf-8") as f:
        spec = yaml.safe_load(f)

    meta = spec.get("meta", {})
    doc = Document()

    # ── Title ──
    title = doc.add_heading(meta.get("name", "Survey Questions"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if meta.get("description"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(meta["description"])
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # ── Consent / Preamble ──
    consent = spec.get("consent", {})
    if consent.get("html"):
        doc.add_heading("Preamble / Consent", level=1)
        text = strip_html(consent["html"])
        for para in text.split("\n"):
            if para.strip():
                p = doc.add_paragraph(para.strip())
                for run in p.runs:
                    run.font.size = Pt(10)

    # ── Build block lookup ──
    blocks_by_id = {}
    for block in spec.get("blocks", []):
        blocks_by_id[block["id"]] = block

    # ── Determine block order from flow ──
    flow_spec = spec.get("flow", {})
    flow_order = flow_spec.get("order", [])
    if flow_order:
        ordered_ids = []
        for step in flow_order:
            if "block" in step:
                ordered_ids.append(step["block"])
    else:
        ordered_ids = [b["id"] for b in spec.get("blocks", [])]

    # Also add any blocks not in the flow (catch-all)
    for b in spec.get("blocks", []):
        if b["id"] not in ordered_ids:
            ordered_ids.append(b["id"])

    # ── Render blocks in order ──
    question_num = 1
    for block_id in ordered_ids:
        block = blocks_by_id.get(block_id)
        if not block:
            continue

        title_text = block.get("title", block_id)

        # Block-level display logic
        block_logic = block.get("display_logic", {}).get("condition", {})
        block_note = ""
        if block_logic:
            ref = block_logic.get("question", "?")
            if "is_not" in block_logic:
                block_note = (
                    f'  [Show if "{ref}" \u2260 "{block_logic["is_not"]}"]'
                )
            elif "selected" in block_logic:
                block_note = (
                    f'  [Show if "{ref}" = "{block_logic["selected"]}"]'
                )

        doc.add_heading(title_text, level=2)
        if block_note:
            p = doc.add_paragraph()
            run = p.add_run(block_note)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

        # Page break note
        p = doc.add_paragraph()
        run = p.add_run("\u2500" * 40 + "  page break  " + "\u2500" * 40)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

        questions = block.get("questions", [])
        for q in questions:
            add_question(doc, q, question_num, title_text)
            # Bipolar/slider groups expand into multiple items but are one logical question
            question_num += 1

    # ── Save ──
    doc.save(args.out)
    print(f"Wrote {args.out} ({question_num - 1} questions)")


if __name__ == "__main__":
    main()
